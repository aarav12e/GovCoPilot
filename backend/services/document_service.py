import os, uuid
import aiofiles
from fastapi import UploadFile, HTTPException
from config import UPLOAD_DIR, MAX_FILE_SIZE_MB

ALLOWED = {".pdf", ".docx", ".txt", ".xlsx"}

async def save_upload(file: UploadFile) -> dict:
    ext = os.path.splitext(file.filename)[1].lower()
    if ext not in ALLOWED:
        raise HTTPException(400, f"Unsupported file type. Allowed: PDF, DOCX, TXT, XLSX")
    content = await file.read()
    size_mb = len(content) / (1024 * 1024)
    if size_mb > MAX_FILE_SIZE_MB:
        raise HTTPException(413, f"File too large ({size_mb:.1f} MB). Max {MAX_FILE_SIZE_MB} MB")
    doc_id = str(uuid.uuid4())
    filename = f"{doc_id}{ext}"
    filepath = os.path.join(UPLOAD_DIR, filename)
    os.makedirs(UPLOAD_DIR, exist_ok=True)
    async with aiofiles.open(filepath, "wb") as f:
        await f.write(content)
    return {"doc_id": doc_id, "filename": file.filename, "filepath": filepath,
            "size_mb": round(size_mb, 2), "extension": ext}

def extract_text(filepath: str, extension: str) -> str:
    ext = extension.lower()
    try:
        if ext == ".pdf":
            from PyPDF2 import PdfReader
            reader = PdfReader(filepath)
            pages = [p.extract_text() for p in reader.pages if p.extract_text()]
            return "\n\n".join(pages) if pages else "Could not extract text from PDF."
        elif ext == ".docx":
            from docx import Document
            doc = Document(filepath)
            return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())
        elif ext == ".txt":
            with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()
        elif ext == ".xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
            lines = []
            for sheet in wb.worksheets:
                lines.append(f"Sheet: {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    row_text = "\t".join(str(c) if c else "" for c in row)
                    if row_text.strip():
                        lines.append(row_text)
            return "\n".join(lines)
    except Exception as e:
        raise HTTPException(422, f"Could not extract text: {str(e)}")
    raise HTTPException(400, f"Unsupported extension: {ext}")
